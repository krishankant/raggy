"""Word, PowerPoint and Excel files, as Markdown with locations.

- DOCX: headings from paragraph styles ("Heading 2" → ##), tables in place.
  Word has no fixed pages, so there are no page labels.
- PPTX: one section per slide ("## Slide 4: Pricing"), its text and tables,
  then its speaker notes; each slide is a location ("slide 4").
- XLSX: one table per sheet with the first non-empty row as the header,
  formula cells as their last computed values; each sheet is a location.
"""

from __future__ import annotations

import io
import zipfile

from raggy.ingest import loaders
from raggy.ingest.images import FigureReader
from raggy.ingest.loaders.tables import markdown_table


def _bad(e: Exception, kind: str) -> loaders.LoadError:
    return loaders.LoadError(f"not a readable {kind} file ({type(e).__name__}: {e})")


def load_docx(data: bytes, options) -> tuple[str, str | None]:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        d = Document(io.BytesIO(data))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as e:
        raise _bad(e, "DOCX") from e
    from docx.oxml.ns import qn

    out = []
    figures, k = FigureReader(options), 0
    for el in d.element.body.iterchildren():  # paragraphs and tables, in document order
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(el, d)
            text = p.text.strip()
            blocks = []
            for blip in el.iter(qn("a:blip")):  # pictures anchored in this paragraph, in order
                part = d.part.related_parts.get(blip.get(qn("r:embed")))
                if part is None:
                    continue  # a linked (not embedded) picture: no bytes to read
                k += 1
                if (block := figures.read(part.blob, f"{figures.name} figure {k}")) is not None:
                    blocks.append(block)
            out += blocks if not text else []
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if style == "Title":
                text = "# " + text
            elif style.startswith("Heading ") and style[8:].isdigit():
                text = "#" * min(int(style[8:]), 6) + " " + text
            out.append(text)
            out += blocks
        elif tag == "tbl":
            t = Table(el, d)
            out.append(markdown_table([[c.text for c in row.cells] for row in t.rows]))
    return "\n\n".join(x for x in out if x), (d.core_properties.title or None)


def load_pptx(data: bytes, options) -> tuple[list[tuple[str, str]], str | None]:
    from pptx import Presentation
    from pptx.exc import PackageNotFoundError

    try:
        prs = Presentation(io.BytesIO(data))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as e:
        raise _bad(e, "PPTX") from e
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    pages, first_title = [], None
    figures = FigureReader(options)
    for n, slide in enumerate(prs.slides, start=1):
        k = 0
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        first_title = first_title or title or None
        parts = [f"## Slide {n}: {title}" if title else f"## Slide {n}"]
        title_id = title_shape.shape_id if title_shape is not None else None
        for shape in slide.shapes:
            if shape.shape_id == title_id:  # python-pptx returns a new proxy per access: compare ids, not objects
                continue
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                k += 1
                try:
                    blob = shape.image.blob
                except (ValueError, KeyError, AttributeError):
                    continue  # an empty picture placeholder or a linked image
                if (block := figures.read(blob, f"{figures.name} slide {n}, picture {k}", f"slide {n}")) is not None:
                    parts.append(block)
            elif getattr(shape, "has_table", False) and shape.has_table:
                parts.append(markdown_table([[c.text for c in row.cells] for row in shape.table.rows]))
            elif shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append("Speaker notes: " + " ".join(slide.notes_slide.notes_text_frame.text.split()))
        pages.append((f"slide {n}", "\n\n".join(parts)))
    return pages, (prs.core_properties.title or first_title)


def load_xlsx(data: bytes, options) -> tuple[list[tuple[str, str]], str | None]:
    import openpyxl
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, ValueError) as e:
        raise _bad(e, "XLSX") from e
    pages = []
    for ws in wb.worksheets:
        rows = [r for r in ws.iter_rows(values_only=True) if any(c not in (None, "") for c in r)]
        if not rows:
            continue
        body = markdown_table(rows) if len(rows) > 1 else " ".join(str(c) for c in rows[0] if c is not None)
        pages.append((f"sheet {ws.title}", f"## Sheet: {ws.title}\n\n{body}"))
    title = wb.properties.title if wb.properties is not None else None
    wb.close()
    return pages, title or None
