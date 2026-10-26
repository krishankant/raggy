from pathlib import Path

import openpyxl
from docx import Document as Docx
from pptx import Presentation
from pptx.util import Inches

from raggy.chunking import ChunkSettings, chunk_document
from raggy.chunking.tokens import simple_count
from raggy.ingest.loaders import load_file
from raggy.ingest.walk import iter_files


def make_docx(path: Path) -> Path:
    d = Docx()
    d.core_properties.title = "Staff rota"
    d.add_heading("Workshop rota", level=1)
    d.add_paragraph("Rota for the Merrow workshop, week of 2 March.")
    d.add_heading("Mechanics", level=2)
    t = d.add_table(rows=1, cols=3)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text, t.rows[0].cells[2].text = "Day", "Mechanic", "Shift"
    for row in (("Tuesday", "Ines", "9-18"), ("Wednesday", "Tomas", "9-18")):
        cells = t.add_row().cells
        for c, v in zip(cells, row):
            c.text = v
    d.save(path)
    return path


def make_pptx(path: Path) -> Path:
    p = Presentation()
    for i, (title, body, notes) in enumerate([("Ebb E-Cargo launch", "Spring 2026", "Welcome everyone."),
                                               ("Pricing", "4,950 credits at launch", "Fleet discounts on slide 3."),
                                               ("Fleet", "10+ bikes: 4,450 credits each", "")], start=1):
        s = p.slides.add_slide(p.slide_layouts[1])
        s.shapes.title.text = title
        s.placeholders[1].text = body
        if notes:
            s.notes_slide.notes_text_frame.text = notes
    p.save(path)
    return path


def make_xlsx(path: Path, rows=3) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stock"
    ws.append(["Code", "Part", "In stock", "Price"])
    for i in range(rows):
        ws.append([f"BW-{i:04d}", f"Part {i}", i % 7, f"=10+{i}"])
    ws.append([None, None, None, None])  # an empty row is skipped
    wb.create_sheet("Notes").append(["Restock every Monday."])
    wb.save(path)
    return path


def test_docx_headings_and_table(tmp_path):
    doc = load_file(make_docx(tmp_path / "rota.docx"))
    assert "# Workshop rota" in doc.text and "## Mechanics" in doc.text
    assert "| Day | Mechanic | Shift |" in doc.text and "| Wednesday | Tomas | 9-18 |" in doc.text
    assert doc.metadata["title"] == "Staff rota" and doc.metadata["type"] == "docx"


def test_pptx_slides_titles_notes_and_labels(tmp_path):
    doc = load_file(make_pptx(tmp_path / "deck.pptx"))
    assert "## Slide 2: Pricing" in doc.text and "4,950 credits at launch" in doc.text
    assert doc.text.count("Pricing") == 1  # the title isn't repeated as body text
    assert "Speaker notes: Fleet discounts on slide 3." in doc.text
    labels = [label for label, _ in doc.metadata["pages"]]
    assert labels == ["slide 1", "slide 2", "slide 3"]
    start = dict(doc.metadata["pages"])["slide 2"]
    assert doc.text[start:].startswith("## Slide 2: Pricing")


def test_xlsx_sheets_as_tables_with_labels(tmp_path):
    doc = load_file(make_xlsx(tmp_path / "stock.xlsx"))
    assert "## Sheet: Stock" in doc.text and "| Code | Part | In stock | Price |" in doc.text
    assert "| BW-0001 | Part 1 | 1 |" in doc.text
    assert [label for label, _ in doc.metadata["pages"]] == ["sheet Stock", "sheet Notes"]
    assert "Restock every Monday." in doc.text and "| | | | |" not in doc.text


def test_xlsx_formulas_give_values(tmp_path):
    f = make_xlsx(tmp_path / "stock.xlsx")
    wb = openpyxl.load_workbook(f)  # openpyxl writes no cached values; simulate a saved-by-Excel file
    ws = wb["Stock"]
    ws["D2"] = 10
    wb.save(f)
    assert "| BW-0000 | Part 0 | 0 | 10 |" in load_file(f).text


def test_csv_sniffs_delimiter(tmp_path):
    f = tmp_path / "deliveries.csv"
    f.write_text("date;city;bikes\n2026-03-01;Port Edda;4\n2026-03-02;Calder;2\n", encoding="utf-8")
    doc = load_file(f)
    assert "| date | city | bikes |" in doc.text and "| 2026-03-02 | Calder | 2 |" in doc.text
    assert doc.metadata["type"] == "csv"


def test_big_sheet_splits_with_header(tmp_path):
    doc = load_file(make_xlsx(tmp_path / "big.xlsx", rows=2000))
    chunks = chunk_document(doc, ChunkSettings(chunker="structure", tokens=200, overlap=0, parent_tokens=0),
                            simple_count)
    rows = [c for c in chunks if "BW-1500" in c.text]
    assert len(chunks) > 20 and rows and "| Code | Part | In stock | Price |" in rows[0].context


def test_walk_includes_office(tmp_path):
    for n in ("a.docx", "b.pptx", "c.xlsx", "d.csv", "e.doc"):
        (tmp_path / n).write_bytes(b"x")
    assert [p.name for p in iter_files(tmp_path)] == ["a.docx", "b.pptx", "c.xlsx", "d.csv"]
