import io
import warnings
from pathlib import Path

import docx
import pytest
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader as RLImage
from reportlab.pdfgen import canvas

from raggy.config import Config
from raggy.embed.fake import FakeEmbedder
from raggy.generate.fake import FakeLLM
from raggy.generate.llm import LLMError
from raggy.ingest.images import ImageError
from raggy.ingest.loaders import LoadError, LoadOptions, load_file
from raggy.pipeline import Raggy


def picture(path: Path, text="Jan 42", size=(400, 300), colour="white") -> Path:
    im = Image.new("RGB", size, colour)
    ImageDraw.Draw(im).text((20, 20), text, fill="black")
    path.parent.mkdir(parents=True, exist_ok=True)
    im.save(path)
    return path


class Reader:
    """Stands in for ImageReader.read: records labels, writes a fixed block."""

    def __init__(self, fail_first=False):
        self.labels, self.fail_first = [], fail_first

    def __call__(self, data: bytes, label: str):
        self.labels.append(label)
        if self.fail_first and len(self.labels) == 1:
            raise ImageError("not a readable image (UnidentifiedImageError)")
        return f"[Image: {label}]\nA described picture."


def opts(reader, **kw):
    return LoadOptions(images="ocr", read_image=reader, **kw)


def make_pdf(path: Path, pages) -> Path:
    """pages: list of lists of ("text", y, s) / ("image", y, png_path)."""
    c = canvas.Canvas(str(path), pagesize=A4)
    for items in pages:
        for kind, y, val in items:
            if kind == "text":
                c.setFont("Helvetica", 10)
                c.drawString(72, y, val)
            else:
                c.drawImage(RLImage(str(val)), 72, y, width=300, height=200)
        c.showPage()
    c.save()
    return path


def test_image_file_loads_block(tmp_path):
    r = Reader()
    doc = load_file(picture(tmp_path / "chart.png"), opts(r))
    assert doc.text == "[Image: chart.png]\nA described picture."
    assert doc.metadata["type"] == "image" and r.labels == ["chart.png"]


def test_image_file_off_mode_is_skipped_with_hint(tmp_path):
    with pytest.raises(LoadError, match="RAGGY_IMAGES=ocr"):
        load_file(picture(tmp_path / "chart.png"), LoadOptions())
    with pytest.raises(LoadError, match="chart.png"):
        load_file(tmp_path / "chart.png", LoadOptions())


def test_decorative_image_file_is_a_load_error(tmp_path):
    with pytest.raises(LoadError, match="too small"):
        load_file(picture(tmp_path / "icon.png", size=(40, 40)), LoadOptions(images="ocr", read_image=lambda d, l: None))


def test_pdf_figure_block_on_its_page(tmp_path):
    fig = picture(tmp_path / "f.png", "Wheelbase 1,850 mm")
    pdf = make_pdf(tmp_path / "spec.pdf", [
        [("text", 760, "Ebb spec sheet. Dimensions are shown in the figure on page 2.")],
        [("text", 760, "Dimensions of the Ebb E-Cargo, in millimetres:"), ("image", 500, fig),
         ("text", 460, "Measured unladen.")],
    ])
    r = Reader()
    doc = load_file(pdf, opts(r))
    assert r.labels == ["spec.pdf p. 2, figure 1"]
    i = doc.text.index("[Image: spec.pdf p. 2, figure 1]")
    assert doc.text.index("Dimensions of the Ebb") < i < doc.text.index("Measured unladen")  # in reading order
    page_at = max(label for label, off in doc.metadata["pages"] if off <= i)
    assert page_at == "p. 2"


def test_corrupt_figure_skipped_document_loads(tmp_path):
    a, b = picture(tmp_path / "a.png", "A"), picture(tmp_path / "b.png", "B", colour="lightblue")
    # (the page needs 20+ characters of text, or it's treated as a scan)
    pdf = make_pdf(tmp_path / "two.pdf", [[("text", 760, "Two figures follow on this page."), ("image", 500, a),
                                           ("image", 250, b)]])
    with warnings.catch_warnings(record=True) as w:
        warnings.simplefilter("always")
        doc = load_file(pdf, opts(Reader(fail_first=True)))
    assert "Two figures follow" in doc.text and "figure 2" in doc.text and "figure 1]" not in doc.text
    assert any("figure 1" in str(x.message) and "skipped" in str(x.message) for x in w)


def test_pptx_picture_block(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Regional team"
    s.shapes.add_picture(str(picture(tmp_path / "org.png", "Lead: Priya Nand")), Inches(1), Inches(2))
    prs.save(tmp_path / "team.pptx")
    r = Reader()
    doc = load_file(tmp_path / "team.pptx", opts(r))
    assert r.labels == ["team.pptx slide 1, picture 1"]
    assert doc.text.index("## Slide 1: Regional team") < doc.text.index("[Image: team.pptx slide 1, picture 1]")


def test_docx_inline_image_in_order(tmp_path):
    d = docx.Document()
    d.add_paragraph("Route for Tuesday deliveries:")
    d.add_picture(str(picture(tmp_path / "map.png", "Stop 1 Arlen Quay")))
    d.add_paragraph("Drivers return by 17:00.")
    d.save(tmp_path / "route.docx")
    r = Reader()
    doc = load_file(tmp_path / "route.docx", opts(r))
    assert r.labels == ["route.docx figure 1"]
    t = doc.text
    assert t.index("Route for Tuesday") < t.index("[Image: route.docx figure 1]") < t.index("Drivers return")


def test_repeated_image_read_once(tmp_path):
    logo = picture(tmp_path / "logo.png", "Brightwater", size=(300, 120))
    prs = Presentation()
    for n in range(3):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = f"Slide {n}"
        s.shapes.add_picture(str(logo), Inches(8), Inches(0.2))
    prs.save(tmp_path / "deck.pptx")
    r = Reader()
    doc = load_file(tmp_path / "deck.pptx", opts(r))
    assert r.labels == ["deck.pptx slide 1, picture 1"] and doc.text.count("[Image:") == 1


def test_scanned_page_not_double_read(tmp_path, monkeypatch):
    scan = picture(tmp_path / "scan.png", "The Fenwick workshop closes on 14 March.", size=(800, 600))
    pdf = make_pdf(tmp_path / "notice.pdf", [[("image", 300, scan)]])  # no text layer: a scan
    monkeypatch.setattr("raggy.ingest.ocr.ocr_image", lambda image: "The Fenwick workshop closes on 14 March.")
    r = Reader()
    doc = load_file(pdf, opts(r, ocr=True))
    assert "Fenwick workshop closes" in doc.text and r.labels == []  # OCR'd as a page, not as a figure


def test_ingest_reports_image_failures(tmp_path, monkeypatch):
    monkeypatch.setattr("raggy.ingest.ocr.ocr_image", lambda image: "Jan 42")
    picture(tmp_path / "docs" / "chart.png")
    c = Config(index_dir=tmp_path / "i", embed_model="fake", llm_provider="fake", llm_model=None, images="vision")
    r = Raggy(c, embedder=FakeEmbedder(), llm=FakeLLM(), vision_llm=FakeLLM(error=LLMError("down")))
    report = r.ingest([str(tmp_path / "docs")])
    assert (report.image_calls, report.image_failed, report.added) == (1, 1, 1)
    hits = r.search("Jan 42")
    assert hits and hits[0].chunk.text.startswith("[Image: chart.png]")  # the OCR text still made it in
