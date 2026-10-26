"""Build the `rich` eval suite's documents: answers that live in tables, slides,
spreadsheets, a CSV and a scanned page (week 9).

    uv run python evals/rich/make_docs.py

Deterministic content; the files are committed. reportlab (dev group) writes the
PDFs; Pillow renders the scanned page; python-docx, python-pptx and openpyxl
write the Office files.
"""

from pathlib import Path

HERE = Path(__file__).parent
OUT = HERE / "docs"


def fleet_pdf(path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Table, TableStyle

    st, grid = getSampleStyleSheet(), TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)])
    flow = [
        Paragraph("Brightwater fleet price list 2026", st["Title"]),
        Paragraph("Ebb E-Cargo fleet pricing", st["Heading2"]),
        Paragraph("Businesses that run delivery fleets get volume pricing on the Ebb E-Cargo. "
                  "Prices are per bike and include assembly at the Fenwick workshop.", st["BodyText"]),
        Table([["Bikes ordered", "Price per bike", "Lead time"], ["1-4", "4,950 credits", "3 weeks"],
               ["5-9", "4,700 credits", "4 weeks"], ["10-24", "4,450 credits", "6 weeks"],
               ["25+", "4,200 credits", "8 weeks"]], style=grid),
        PageBreak(),
        Paragraph("Fleet service plans", st["Heading2"]),
        Paragraph("Every fleet bike can be covered by a service plan, billed yearly per bike.", st["BodyText"]),
        Table([["Plan", "Workshop visits per year", "Price per bike per year"], ["Basic", "1", "300 credits"],
               ["Plus", "2", "520 credits"], ["Total care", "4", "900 credits"]], style=grid),
        PageBreak(),
        Paragraph("Terms", st["Heading2"]),
        Paragraph("Fleet prices are valid until 31 December 2026. Orders cancelled after dispatch pay a "
                  "handling fee of 5% of the order value. Fleet invoices are payable within 30 days.", st["BodyText"]),
    ]
    SimpleDocTemplate(str(path), pagesize=A4, title="Brightwater fleet price list 2026",
                      invariant=1).build(flow)


def scanned_pdf(path: Path) -> None:
    from PIL import Image, ImageDraw, ImageFont

    lines = ["WORKSHOP NOTICE", "The Fenwick workshop closes for stocktaking on 14 March.",
             "Collections move to Merrow that day.", "Repairs booked for 14 March move to 15 March."]
    img = Image.new("L", (1700, 420), 255)
    d = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 40)
    except OSError:
        font = ImageFont.load_default(size=40)
    for i, line in enumerate(lines):
        d.text((50, 40 + i * 90), line, fill=0, font=font)
    img.convert("RGB").save(path, "PDF", resolution=150)


def rota_docx(path: Path) -> None:
    from docx import Document

    d = Document()
    d.core_properties.title = "Workshop rota"
    d.add_heading("Merrow workshop rota", level=1)
    d.add_paragraph("Week of 2 March. Saturday shifts rotate between Ines and Tomas.")
    d.add_heading("Mechanics on shift", level=2)
    rows = [("Day", "Mechanic", "Shift"), ("Tuesday", "Ines Okafor", "9:00-18:00"),
            ("Wednesday", "Tomas Rehn", "9:00-18:00"), ("Thursday", "Priya Nand", "12:00-21:00"),
            ("Friday", "Ines Okafor", "9:00-18:00"), ("Saturday", "Tomas Rehn", "9:00-16:00")]
    t = d.add_table(rows=0, cols=3)
    for r in rows:
        cells = t.add_row().cells
        for c, v in zip(cells, r):
            c.text = v
    d.add_paragraph("Late Thursday shifts cover the community repair evening.")
    d.save(path)


def launch_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    slides = [("Ebb E-Cargo 2026 launch", "Launch day: 21 April 2026", "Press embargo lifts at 09:00 on launch day."),
              ("Price", "4,950 credits at launch, including assembly", ""),
              ("Range", "Up to 90 km on Eco with no cargo", "Range drops to about 70 km with 80 kg in the box."),
              ("Colours", "Harbour blue, Sand and Charcoal", ""),
              ("Fleet pricing", None, "Fleet price list has the full table."),
              ("Next steps", "Dealer training in the week of 7 April", "")]
    prs = Presentation()
    for title, body, notes in slides:
        s = prs.slides.add_slide(prs.slide_layouts[5 if body is None else 1])
        s.shapes.title.text = title
        if body is None:
            rows = [("Bikes", "Price per bike"), ("1-4", "4,950 credits"), ("10-24", "4,450 credits")]
            t = s.shapes.add_table(len(rows), 2, Inches(1), Inches(2), Inches(6), Inches(2)).table
            for i, r in enumerate(rows):
                for j, v in enumerate(r):
                    t.cell(i, j).text = v
        else:
            s.placeholders[1].text = body
        if notes:
            s.notes_slide.notes_text_frame.text = notes
    prs.core_properties.title = "Ebb E-Cargo 2026 launch"
    prs.save(path)


def stock_xlsx(path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stock"
    ws.append(["Code", "Part", "Location", "In stock"])
    special = {7: ("BW-CH-1142", "8-speed chain, 116 links", "Merrow", 42),
               150: ("BW-BT-5000", "500 Wh replacement battery", "Fenwick", 6),
               290: ("XP-TY-4545", "45 mm gravel tyre", "Merrow", 18)}
    for i in range(300):
        ws.append(list(special.get(i, (f"BW-GN-{i:04d}", f"Generic part {i}", ["Merrow", "Fenwick"][i % 2],
                                        (i * 7) % 50))))
    sup = wb.create_sheet("Suppliers")
    for r in [("Supplier", "Parts", "Lead time"), ("Kessel Components", "Chains and cassettes", "2 weeks"),
              ("Northline Cells", "Batteries", "5 weeks"), ("Tarn Tyres", "Tyres and tubes", "10 days")]:
        sup.append(r)
    wb.properties.title = "Parts stock"
    wb.save(path)


def deliveries_csv(path: Path) -> None:
    rows = ["date;city;bikes;driver"]
    cities = ["Port Edda", "Calder", "Arlen Islands", "Merrow"]
    for d in range(1, 21):
        rows.append(f"2026-03-{d:02d};{cities[d % 4]};{(d * 3) % 7 + 1};{['Sam', 'Ola', 'Rui'][d % 3]}")
    path.write_text("\n".join(rows) + "\n", encoding="utf-8")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    fleet_pdf(OUT / "fleet-prices.pdf")
    scanned_pdf(OUT / "workshop-notice-scan.pdf")
    rota_docx(OUT / "workshop-rota.docx")
    launch_pptx(OUT / "ebb-launch.pptx")
    stock_xlsx(OUT / "parts-stock.xlsx")
    deliveries_csv(OUT / "deliveries.csv")
    print("wrote", sorted(p.name for p in OUT.iterdir()))


if __name__ == "__main__":
    main()
