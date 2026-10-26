"""Build the `images` eval suite's documents: answers that exist only in pictures (week 10).

    uv run python evals/images/make_docs.py

Deterministic content; the files are committed. Pillow draws every picture;
reportlab, python-docx and python-pptx put some of them inside documents.
Each picture is built for one kind of question:

- text-in-image: the answer is printed in the picture (OCR should be enough)
- association: OCR reads the parts but loses which value belongs to which label
- visual: no words at all (a colour, a shape, a count): only a vision model can answer
"""

import io
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

HERE = Path(__file__).parent
OUT = HERE / "docs"


def font(size: int):
    try:
        return ImageFont.truetype("Helvetica", size)
    except OSError:
        return ImageFont.load_default(size)


def service_chart() -> Image.Image:
    """Bar chart: month → bikes serviced (association)."""
    im = Image.new("RGB", (900, 560), "white")
    d = ImageDraw.Draw(im)
    d.text((40, 20), "Bikes serviced per month, Merrow workshop 2026", fill="black", font=font(28))
    for i, (m, v) in enumerate([("Jan", 42), ("Feb", 57), ("Mar", 81), ("Apr", 64), ("May", 73)]):
        x = 80 + i * 160
        d.rectangle([x, 500 - v * 4.5, x + 90, 500], fill=(40, 90, 200))
        d.text((x + 28, 510), m, fill="black", font=font(24))
        d.text((x + 28, 470 - v * 4.5), str(v), fill="black", font=font(24))
    return im


def dimension_drawing() -> Image.Image:
    """Side view with dimension lines (association)."""
    im = Image.new("RGB", (1000, 560), "white")
    d = ImageDraw.Draw(im)
    d.text((40, 20), "Ebb E-Cargo dimensions (mm)", fill="black", font=font(28))
    d.ellipse([120, 330, 280, 490], outline="black", width=6)
    d.ellipse([720, 330, 880, 490], outline="black", width=6)
    d.rectangle([150, 220, 460, 330], outline="black", width=5)  # the cargo box
    d.line([200, 520, 800, 520], fill="black", width=3)
    d.text((400, 525), "Wheelbase 1,850", fill="black", font=font(24))
    d.line([480, 100, 480, 490], fill="black", width=3)
    d.text((495, 150), "Height 1,120", fill="black", font=font(24))
    d.text((160, 180), "Box: 80 kg max", fill="black", font=font(24))
    return im


def floorplan() -> Image.Image:
    """Labelled rooms and an arrow to the fire exit (association)."""
    im = Image.new("RGB", (1000, 640), "white")
    d = ImageDraw.Draw(im)
    d.text((40, 20), "Merrow workshop floor plan", fill="black", font=font(28))
    rooms = [("Reception", 40, 80, 340, 300), ("Wash bay", 340, 80, 640, 300), ("Paint booth", 640, 80, 960, 300),
             ("Bay 3: battery charging", 40, 300, 500, 600), ("Parts store", 500, 300, 960, 600)]
    for name, x0, y0, x1, y1 in rooms:
        d.rectangle([x0, y0, x1, y1], outline="black", width=4)
        d.text((x0 + 20, y0 + 20), name, fill="black", font=font(24))
    d.line([800, 560, 960, 560], fill=(200, 30, 30), width=8)
    d.polygon([(960, 545), (985, 560), (960, 575)], fill=(200, 30, 30))
    d.text((760, 520), "Fire exit", fill=(200, 30, 30), font=font(24))
    return im


def battery_label() -> Image.Image:
    """A rating label (text-in-image)."""
    im = Image.new("RGB", (900, 420), (235, 235, 225))
    d = ImageDraw.Draw(im)
    d.rectangle([20, 20, 880, 400], outline="black", width=4)
    for i, line in enumerate(["BRIGHTWATER BW-BT-5000", "Li-ion 36 V  13.9 Ah  500 Wh",
                              "Charge at 0-40 °C only", "Assembled in Fenwick"]):
        d.text((50, 50 + i * 85), line, fill="black", font=font(36))
    return im


def torque_table() -> Image.Image:
    """A screenshot of a table (association)."""
    rows = [("Part", "Torque"), ("Stem bolts", "5 Nm"), ("Handlebar clamp", "6 Nm"), ("Pedals", "35 Nm"),
            ("Rear axle nut", "40 Nm")]
    im = Image.new("RGB", (760, 120 + 70 * len(rows)), "white")
    d = ImageDraw.Draw(im)
    d.text((30, 20), "Ebb torque settings", fill="black", font=font(28))
    for i, (a, b) in enumerate(rows):
        y = 80 + i * 70
        d.rectangle([30, y, 730, y + 70], outline="black", width=2)
        d.line([480, y, 480, y + 70], fill="black", width=2)
        d.text((50, y + 20), a, fill="black", font=font(26))
        d.text((500, y + 20), b, fill="black", font=font(26))
    return im


def product_drawing() -> Image.Image:
    """A bike with a blue cargo box and two child seats, no text at all (visual)."""
    im = Image.new("RGB", (900, 520), (245, 245, 240))
    d = ImageDraw.Draw(im)
    d.ellipse([100, 300, 280, 480], outline="black", width=8)
    d.ellipse([640, 300, 820, 480], outline="black", width=8)
    d.rectangle([110, 160, 420, 300], fill=(30, 80, 200), outline="black", width=4)
    for x in (150, 290):  # two child seats in the box
        d.rectangle([x, 110, x + 90, 160], fill=(240, 200, 40), outline="black", width=3)
    d.line([420, 300, 730, 390], fill="black", width=8)
    d.line([600, 180, 730, 390], fill="black", width=8)
    d.line([580, 170, 640, 170], fill="black", width=8)
    return im


def org_chart() -> Image.Image:
    """Boxes and lines: who reports to whom (association)."""
    im = Image.new("RGB", (1000, 520), "white")
    d = ImageDraw.Draw(im)
    d.text((40, 20), "Regional service team", fill="black", font=font(28))
    boxes = {"Head of service: Marta Lindqvist": (300, 80, 700, 160), "Merrow lead: Priya Nand": (60, 300, 440, 380),
             "Fenwick lead: Joel Abara": (560, 300, 940, 380)}
    for name, (x0, y0, x1, y1) in boxes.items():
        d.rectangle([x0, y0, x1, y1], outline="black", width=4)
        d.text((x0 + 15, y0 + 25), name, fill="black", font=font(22))
    d.line([500, 160, 500, 230, 250, 230, 250, 300], fill="black", width=3)
    d.line([500, 230, 750, 230, 750, 300], fill="black", width=3)
    return im


def route_map() -> Image.Image:
    """Numbered stops along a route (association)."""
    im = Image.new("RGB", (1000, 600), (240, 245, 240))
    d = ImageDraw.Draw(im)
    d.text((40, 20), "Tuesday delivery route", fill="black", font=font(28))
    stops = [("Merrow depot", 100, 480), ("1 Arlen Quay", 300, 300), ("2 Calder Market", 560, 380),
             ("3 Harbour Lane", 800, 200)]
    d.line([(x, y) for _, x, y in stops], fill=(200, 60, 30), width=6)
    for name, x, y in stops:
        d.ellipse([x - 14, y - 14, x + 14, y + 14], fill="black")
        d.text((x + 20, y - 40), name, fill="black", font=font(24))
    return im


def logo() -> Image.Image:
    im = Image.new("RGB", (64, 64), (30, 80, 200))  # decorative: under 100 px, skipped
    ImageDraw.Draw(im).text((10, 20), "BW", fill="white", font=font(24))
    return im


def png_bytes(im: Image.Image) -> bytes:
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue()


def spec_pdf(path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A4)
    c.setTitle("Ebb E-Cargo spec sheet")
    c.setFont("Helvetica-Bold", 18)
    c.drawString(72, 770, "Ebb E-Cargo spec sheet")
    c.setFont("Helvetica", 11)
    c.drawString(72, 740, "The Ebb E-Cargo is Brightwater's family cargo bike. Dimensions are on page 2.")
    c.showPage()
    c.setFont("Helvetica", 11)
    c.drawString(72, 770, "Dimensions of the Ebb E-Cargo, measured unladen:")
    c.drawImage(ImageReader(io.BytesIO(png_bytes(dimension_drawing()))), 72, 420, width=450, height=252)
    c.drawString(72, 390, "Measured at the Fenwick workshop.")
    c.showPage()
    c.save()


def team_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for title, picture in (("Our regional team", org_chart()), ("Questions?", None)):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = title
        s.shapes.add_picture(io.BytesIO(png_bytes(logo())), Inches(9), Inches(0.1))
        if picture is not None:
            s.shapes.add_picture(io.BytesIO(png_bytes(picture)), Inches(1), Inches(1.8), width=Inches(8))
    prs.save(path)


def route_docx(path: Path) -> None:
    import docx
    from docx.shared import Inches

    d = docx.Document()
    d.core_properties.title = "Tuesday deliveries"
    d.add_heading("Tuesday deliveries", 1)
    d.add_paragraph("The route for Tuesday deliveries is shown below.")
    d.add_picture(io.BytesIO(png_bytes(route_map())), width=Inches(6))
    d.add_paragraph("Drivers return to the depot by 17:00.")
    d.save(path)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    service_chart().save(OUT / "merrow-service-chart.png")
    floorplan().save(OUT / "workshop-floorplan.png")
    battery_label().save(OUT / "bt-5000-label.jpg", quality=92)
    torque_table().save(OUT / "torque-settings.png")
    product_drawing().save(OUT / "ebb-product.png")
    spec_pdf(OUT / "ebb-spec-sheet.pdf")
    team_pptx(OUT / "regional-team.pptx")
    route_docx(OUT / "delivery-route.docx")
    print(f"wrote {len(list(OUT.iterdir()))} files to {OUT}")


if __name__ == "__main__":
    main()
